"""
批量摘要路由：
- POST /api/batch/upload   上传多个 txt 文件，后台逐篇生成摘要
- GET  /api/batch/status/{batch_id}  轮询处理进度与结果
- GET  /api/batch/export/{batch_id}  导出结果为 Excel（openpyxl）
"""
import json
import os
import threading
import time
import uuid
from typing import Dict, List, Optional

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from loguru import logger

from ..model_manager import get_model_manager
from .. import text_tools

router = APIRouter(prefix="/api/batch", tags=["批量处理"])

_BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_BATCH_DIR = os.path.join(_BASE_DIR, "data", "batch")

# 内存中的批处理状态 {batch_id: {...}}；结果同时持久化到 JSON 文件
_batches: Dict[str, Dict] = {}
_lock = threading.Lock()


@router.post("/upload", summary="上传多个 txt 新闻文件或单个 csv 文件开始批量摘要")
async def upload(files: List[UploadFile] = File(...),
                 model: Optional[str] = Form(None),
                 max_length: int = Form(128),
                 min_length: int = Form(32)):
    if not files:
        raise HTTPException(status_code=400, detail="未收到任何文件")
    batch_id = uuid.uuid4().hex[:12]
    batch_dir = os.path.join(_BATCH_DIR, batch_id)
    os.makedirs(batch_dir, exist_ok=True)

    docs = []
    for f in files:
        if not f.filename:
            continue
        lower = f.filename.lower()
        raw = await f.read()
        if len(raw) > 2 * 1024 * 1024:  # 单文件上限 2MB
            raw = raw[:2 * 1024 * 1024]
        safe_name = os.path.basename(f.filename)

        if lower.endswith(".csv"):
            # CSV 导入：每行一条新闻，取 text/正文/content 列或最后一列
            import csv, io as _io
            text = _decode(raw)
            reader = csv.DictReader(_io.StringIO(text))
            rows = list(reader)
            if not rows:
                continue
            fieldnames = reader.fieldnames or []
            text_col = next(
                (fn for fn in fieldnames if fn.strip().lower() in ("text", "正文", "content", "news")),
                None,
            )
            if text_col is None:
                text_col = fieldnames[-1]  # 兜底取最后一列
            for i, row in enumerate(rows):
                body = (row.get(text_col) or "").strip()
                if body:
                    docs.append({"filename": f"{safe_name}#row{i+1}", "text": body})
            continue

        if not lower.endswith(".txt"):
            logger.warning(f"跳过不支持的文件类型: {f.filename}")
            continue
        text = _decode(raw)
        if not text.strip():
            continue
        with open(os.path.join(batch_dir, f"{len(docs)}.txt"), "w", encoding="utf-8") as fp:
            fp.write(text)
        docs.append({"filename": safe_name, "text": text})

    if not docs:
        raise HTTPException(status_code=400, detail="没有可处理的文本（请上传非空 txt 或含 text/正文 列的 csv）")

    state = {
        "batch_id": batch_id,
        "status": "processing",
        "progress": 0,
        "total": len(docs),
        "model": model or "",
        "max_length": int(max_length),
        "min_length": int(min_length),
        "results": [],
        "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "_docs": docs,
    }
    with _lock:
        _batches[batch_id] = state
    # 后台线程逐篇处理，避免请求超时
    threading.Thread(target=_process_batch, args=(batch_id,), daemon=True).start()
    return {"batch_id": batch_id, "total": len(docs), "status": "processing"}


@router.get("/status/{batch_id}", summary="查询批量处理进度与结果")
async def status(batch_id: str):
    state = _get_state(batch_id)
    return {
        "batch_id": batch_id,
        "status": state["status"],
        "progress": state["progress"],
        "total": state["total"],
        "results": state["results"] if state["status"] == "done" else [],
        "error": state.get("error", ""),
    }


@router.get("/export/{batch_id}", summary="批量摘要结果导出为 Excel 或 CSV")
async def export(batch_id: str, format: str = "xlsx"):
    state = _get_state(batch_id)
    if state["status"] != "done" or not state["results"]:
        raise HTTPException(status_code=400, detail="批处理尚未完成，无法导出")

    headers = ["序号", "文件名", "原文字数", "摘要字数", "压缩比(%)",
               "新闻分类", "分类置信度(%)", "可信度(%)", "润色后摘要", "生成模型"]

    def _rows():
        return [
            [i, r["filename"], r["original_chars"], r["summary_chars"],
             r["compression_ratio"], r["category"], r["category_confidence"],
             r["credibility"], r["polished_summary"], r["model_display_name"]]
            for i, r in enumerate(state["results"], 1)
        ]

    if format == "csv":
        import csv, io as _io
        buf = _io.StringIO()
        buf.write("﻿")  # UTF-8 BOM，Excel 打开不乱码
        writer = csv.writer(buf)
        writer.writerow(headers)
        writer.writerows(_rows())
        content = buf.getvalue().encode("utf-8")
        return StreamingResponse(
            _io.BytesIO(content),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="批量摘要结果_{batch_id}.csv"'},
        )

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail="服务器缺少 openpyxl，无法导出 Excel")

    wb = Workbook()
    ws = wb.active
    ws.title = "批量摘要结果"
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="409EFF")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for i, r in enumerate(state["results"], 1):
        ws.append([
            i, r["filename"], r["original_chars"], r["summary_chars"],
            r["compression_ratio"], r["category"], r["category_confidence"],
            r["credibility"], r["polished_summary"], r["model_display_name"],
        ])

    widths = [6, 24, 10, 10, 10, 10, 14, 10, 70, 26]
    from openpyxl.utils import get_column_letter
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    out_path = os.path.join(_BATCH_DIR, batch_id, "result.xlsx")
    wb.save(out_path)
    logger.info(f"批量结果导出: {out_path}")
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"批量摘要结果_{batch_id}.xlsx")


# ---------------------------------------------------------------------------
# 内部工具
# ---------------------------------------------------------------------------
def _decode(raw: bytes) -> str:
    """尝试多种编码解码 txt 内容。"""
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", errors="ignore")


def _get_state(batch_id: str) -> Dict:
    with _lock:
        if batch_id in _batches:
            return _batches[batch_id]
    # 服务重启后从磁盘恢复
    state_file = os.path.join(_BATCH_DIR, batch_id, "state.json")
    if os.path.isfile(state_file):
        with open(state_file, "r", encoding="utf-8") as f:
            state = json.load(f)
        with _lock:
            _batches[batch_id] = state
        return state
    raise HTTPException(status_code=404, detail="批次不存在")


def _save_state(batch_id: str, state: Dict) -> None:
    state_file = os.path.join(_BATCH_DIR, batch_id, "state.json")
    try:
        with open(state_file, "w", encoding="utf-8") as f:
            json.dump(state, f, ensure_ascii=False)
    except Exception as e:
        logger.error(f"保存批次状态失败: {e}")


def _process_batch(batch_id: str) -> None:
    with _lock:
        state = _batches.get(batch_id)
    if state is None:
        return
    manager = get_model_manager()
    model_name = state["model"] or None
    for idx, doc in enumerate(list(state.get("_docs", []))):
        try:
            text = doc["text"][:5000]
            gen = manager.summarize(
                text=text, model_name=model_name,
                max_length=state["max_length"], min_length=state["min_length"])
            polished = text_tools.polish_summary(gen["summary"])["polished"] or gen["summary"]
            stats = text_tools.count_stats(doc["text"], polished)
            cls = text_tools.classify_news(doc["text"])
            cred = text_tools.credibility_score(doc["text"], polished)
            state["results"].append({
                "filename": doc["filename"],
                "original_chars": stats["original_chars"],
                "summary_chars": stats["summary_chars"],
                "compression_ratio": stats["compression_ratio"],
                "category": cls["category"],
                "category_confidence": cls["confidence"],
                "credibility": cred["score"],
                "summary": gen["summary"],
                "polished_summary": polished,
                "model_display_name": gen["model_display_name"],
                "status": "成功",
            })
        except Exception as e:
            logger.error(f"批量处理单篇失败: {doc['filename']}, {e}")
            state["results"].append({
                "filename": doc["filename"], "status": f"失败: {e}",
                "original_chars": 0, "summary_chars": 0, "compression_ratio": 0,
                "category": "", "category_confidence": 0, "credibility": 0,
                "summary": "", "polished_summary": "", "model_display_name": "",
            })
        state["progress"] = idx + 1
        _save_state(batch_id, state)

    state.pop("_docs", None)
    state["status"] = "done"
    _save_state(batch_id, state)
    logger.info(f"批量处理完成: {batch_id}, 共 {state['total']} 篇")
